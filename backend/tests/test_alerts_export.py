from io import BytesIO

from openpyxl import load_workbook

from app.domain.alerts_export import build_alert_excel, flujo_estado_label


def test_flujo_estado_label_espeja_frontend():
    assert flujo_estado_label("ENVIADO", "GERENTE") == "Por validar"
    assert flujo_estado_label("ENVIADO", "ADMIN") == "Pendiente de gerente"
    assert flujo_estado_label("VALIDADO", "ADMIN") == "Por recepcionar"
    assert flujo_estado_label("RECEPCIONADO", "JEFE") == "Recepcionado"
    assert flujo_estado_label(None, "JEFE") == "Borrador"


def _persona(**kw):
    base = {
        "dni": "10000001",
        "nombre": "María Pérez",
        "area": "T.I.",
        "jefatura": "T.I.",
        "jefe_nombre": "Carlos Ríos",
        "gerencia": "Planificación",
        "total_dias": 15,
        "tope_dias": 30,
        "estado_plan": "Parcial (15/30)",
        "flujo_estado": "ENVIADO",
    }
    base.update(kw)
    return base


def test_build_alert_excel_pendientes_flujo_incluye_estado_flujo():
    data = build_alert_excel("pendientes_flujo", [_persona()], "GERENTE")
    wb = load_workbook(BytesIO(data))
    ws = wb["PENDIENTES_FLUJO"]
    headers = [c.value for c in ws[1]]
    assert "ESTADO_FLUJO" in headers
    row = {h: ws.cell(2, i + 1).value for i, h in enumerate(headers)}
    assert row["NOMBRE"] == "María Pérez"
    assert row["JEFE"] == "Carlos Ríos"
    assert row["ESTADO_FLUJO"] == "Por validar"


def test_build_alert_excel_record_incluye_fecha_y_dias():
    persona = _persona(fecha_vencimiento="2026-12-08", dias_restantes=45)
    data = build_alert_excel("record_vence", [persona], "JEFE")
    wb = load_workbook(BytesIO(data))
    ws = wb["RECORD_POR_VENCER"]
    headers = [c.value for c in ws[1]]
    assert "FECHA_VENCIMIENTO" in headers
    assert "DIAS_RESTANTES" in headers
    assert "ESTADO_FLUJO" not in headers


def test_build_alert_excel_mes_siguiente_incluye_periodos():
    persona = _persona(mes="2026-09", dias_mes=5, periodos_mes="14/09–18/09")
    data = build_alert_excel("mes_siguiente", [persona], "GERENTE")
    wb = load_workbook(BytesIO(data))
    ws = wb["MES_SIGUIENTE"]
    headers = [c.value for c in ws[1]]
    assert "DIAS_EN_MES" in headers
    assert "PERIODOS" in headers


def test_build_alert_excel_vacio_no_revienta():
    data = build_alert_excel("sin_programar", [], "JEFE")
    wb = load_workbook(BytesIO(data))
    assert "APTOS_SIN_PROGRAMAR" in wb.sheetnames
