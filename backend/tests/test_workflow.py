from datetime import date

from app.domain.workflow import (
    BORRADOR,
    ENVIADO,
    OBSERVADO,
    RECEPCIONADO,
    VALIDADO,
    allowed_transition,
    apply_transition,
    can_edit_days,
    flujo_from_row,
    reject_if_cannot_edit,
)


def test_jefe_edita_borrador_apto():
    user = {"rol": "JEFE", "is_admin": False}
    emp = {"nombre": "Ana", "fecha_ingreso": "2020-01-15"}
    assert can_edit_days(user, emp, BORRADOR, today=date(2026, 9, 4))
    reject_if_cannot_edit(user, emp, BORRADOR, today=date(2026, 9, 4))


def test_jefe_no_edita_no_apto():
    user = {"rol": "JEFE", "is_admin": False}
    emp = {"nombre": "Nuevo", "fecha_ingreso": "2026-03-01"}
    assert not can_edit_days(user, emp, BORRADOR, today=date(2026, 9, 4))
    try:
        reject_if_cannot_edit(user, emp, BORRADOR, today=date(2026, 9, 4))
    except ValueError as exc:
        assert "cumple el año" in str(exc)
    else:
        raise AssertionError("debía rechazar")


def test_jefe_no_edita_enviado():
    user = {"rol": "JEFE"}
    emp = {"nombre": "Ana", "fecha_ingreso": "2020-01-15"}
    assert not can_edit_days(user, emp, ENVIADO, today=date(2026, 9, 4))


def test_gerente_no_programa():
    user = {"rol": "GERENTE"}
    emp = {"nombre": "Ana", "fecha_ingreso": "2020-01-15"}
    assert not can_edit_days(user, emp, BORRADOR, today=date(2026, 9, 4))


def test_admin_si_edita_recepcionado():
    user = {"rol": "ADMIN", "is_admin": True}
    emp = {"nombre": "Ana", "fecha_ingreso": "2020-01-15"}
    assert can_edit_days(user, emp, RECEPCIONADO, today=date(2026, 9, 4))


def test_transiciones():
    assert allowed_transition("JEFE", BORRADOR, ENVIADO)
    assert allowed_transition("JEFE", OBSERVADO, ENVIADO)
    assert not allowed_transition("JEFE", ENVIADO, VALIDADO)
    assert allowed_transition("GERENTE", ENVIADO, VALIDADO)
    assert allowed_transition("GERENTE", ENVIADO, OBSERVADO)
    assert allowed_transition("ADMIN", VALIDADO, RECEPCIONADO)
    assert allowed_transition("ADMIN", VALIDADO, OBSERVADO)
    assert not allowed_transition("ADMIN", ENVIADO, RECEPCIONADO)


def test_observar_exige_motivo():
    row = flujo_from_row(None)
    row["estado"] = ENVIADO
    user = {"rol": "GERENTE", "correo": "g@x.pe"}
    try:
        apply_transition(row, user, OBSERVADO, observacion="")
    except ValueError as exc:
        assert "motivo" in str(exc).lower()
    else:
        raise AssertionError("debía exigir observación")
    out = apply_transition(row, user, OBSERVADO, observacion="Falta el bloque de 15 días")
    assert out["estado"] == OBSERVADO
    assert "15" in out["observacion"]


def test_excel_flujo_roundtrip():
    from app.domain.workflow_excel import build_flujo_excel, parse_flujo_upload

    employees = [
        {
            "dni": "12345678",
            "nombre": "Ana Perez",
            "area": "T.I.",
            "gerencia": "Planificación",
            "jefatura": "T.I.",
            "tipo_personal": "EMPLEADO",
            "fecha_ingreso": "2020-01-15",
            "apto": True,
            "record_cumplido": True,
            "flujo_estado": "BORRADOR",
            "flujo_observacion": "",
            "cumple_record": "2021-01-15",
            "jefe_correo": "jefe@x.pe",
            "gerente_correo": "",
            "admin_correo": "",
        }
    ]
    daily = {"12345678|2026-09-14"}
    data = build_flujo_excel(employees, daily, {("12345678", 38): 1}, 2026, 2026, 36, rol="JEFE")
    from io import BytesIO
    from openpyxl import load_workbook
    assert load_workbook(BytesIO(data)).sheetnames == ["PERIODOS"]
    flujo, periods = parse_flujo_upload(data)
    assert flujo == []
    assert periods[0]["dni"] == "12345678"
    assert periods[0]["fecha_inicio"].isoformat() == "2026-09-14"


def test_excel_gerente_solo_flujo():
    from io import BytesIO

    from openpyxl import load_workbook

    from app.domain.workflow_excel import build_flujo_excel, parse_flujo_upload

    employees = [
        {
            "dni": "12345678",
            "nombre": "Ana Perez",
            "area": "T.I.",
            "gerencia": "Planificación",
            "jefatura": "T.I.",
            "tipo_personal": "EMPLEADO",
            "fecha_ingreso": "2020-01-15",
            "apto": True,
            "record_cumplido": True,
            "flujo_estado": "ENVIADO",
            "flujo_observacion": "",
            "cumple_record": "2021-01-15",
        }
    ]
    data = build_flujo_excel(employees, {"12345678|2026-09-14"}, {}, 2026, 2026, 36, rol="GERENTE")
    assert load_workbook(BytesIO(data)).sheetnames == ["FLUJO"]
    flujo, periods = parse_flujo_upload(data)
    assert flujo[0]["estado"] == "ENVIADO"
    assert periods == []


def test_excel_no_apto_no_entra_en_flujo():
    from app.domain.workflow_excel import build_flujo_excel, parse_flujo_upload

    no_apto = {
        "dni": "999",
        "nombre": "Nuevo",
        "area": "T.I.",
        "gerencia": "Planificación",
        "jefatura": "T.I.",
        "tipo_personal": "EMPLEADO",
        "fecha_ingreso": "2026-03-01",
        "apto": False,
        "record_cumplido": False,
        "flujo_estado": "BORRADOR",
        "flujo_observacion": "",
        "cumple_record": "2027-03-01",
    }
    data = build_flujo_excel([no_apto], {"999|2026-09-14"}, {}, 2026, 2026, 36, rol="JEFE")
    flujo, periods = parse_flujo_upload(data)
    assert flujo == []
    assert periods == []


def test_excel_rechaza_mas_de_4mb():
    import asyncio

    from app.upload_limit import UploadTooLarge, read_upload_limited

    class FakeFile:
        def __init__(self):
            self.n = 0

        async def read(self, _n):
            self.n += 1
            if self.n > 70:
                return b""
            return b"x" * 64 * 1024

    try:
        asyncio.run(read_upload_limited(FakeFile()))
    except UploadTooLarge:
        return
    raise AssertionError("debía rechazar un Excel de más de 4 MB")


def test_excel_periodos_rechaza_pasar_del_vencimiento():
    """Pepe (récord mar-2026/mar-2027, goce hasta mar-2028): el Excel no debe aceptar abril-2028."""
    from app.domain.workflow import apply_uploaded_periods

    ingreso = date(2026, 3, 15)
    daily: set[str] = set()
    targets: dict = {}
    periods = [{"fecha_inicio": date(2028, 4, 1), "fecha_fin": date(2028, 4, 5)}]
    try:
        apply_uploaded_periods(
            daily,
            targets,
            "1",
            2026,
            periods,
            today=date(2026, 9, 7),
            fecha_ingreso=ingreso,
            nombre="Pepe",
        )
    except ValueError as exc:
        assert "14/03/2028" in str(exc)
    else:
        raise AssertionError("debía rechazar fechas después del vencimiento del récord")
    assert not daily


def test_excel_periodos_respeta_tope_de_dias_por_fila():
    """Un rango de fechas absurdamente largo en una fila del Excel debe rechazarse antes de aplicarse."""
    from io import BytesIO

    import pandas as pd

    from app.domain.workflow_excel import parse_flujo_upload

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(
            [{"DNI": "12345678", "FECHA_INICIO": "01/01/2020", "FECHA_FIN": "01/01/2030"}]
        ).to_excel(writer, sheet_name="PERIODOS", index=False)
    try:
        parse_flujo_upload(output.getvalue())
    except ValueError as exc:
        assert "máximo" in str(exc)
    else:
        raise AssertionError("debía rechazar un rango de fechas demasiado largo")


def test_excel_estado_invalido():
    from io import BytesIO

    import pandas as pd

    from app.domain.workflow_excel import parse_flujo_upload

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame([{"DNI": "1", "ESTADO": "HACKED"}]).to_excel(writer, sheet_name="FLUJO", index=False)
    try:
        parse_flujo_upload(output.getvalue())
    except ValueError as exc:
        assert "ESTADO no válido" in str(exc)
    else:
        raise AssertionError("debía rechazar estado inválido")
